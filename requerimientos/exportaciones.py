"""Exportación de la bandeja de compras a Excel (HU-26).

«Como analista de compras quiero exportar los requerimientos a Excel» para
alimentar reportes de área sin transcribirlos a mano desde la pantalla. El
archivo recoge exactamente lo que el analista tiene filtrado en la bandeja
(HU-21): si filtró por área y estado, el Excel trae solo eso, no todo el
sistema — igual que al exportar desde cualquier listado con filtros.

Se usa `openpyxl` directamente y no el motor de HU-30 (que vuelca vistas SQL
a CSV para Power BI): esa exportación es para analítica agregada, esta es el
detalle operativo de la bandeja, en el formato que el área de compras abre
todos los días.
"""

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ENCABEZADOS = (
    "Consecutivo",
    "Estado",
    "Fecha de solicitud",
    "Solicitante",
    "Área o proyecto",
    "Centro de costo",
    "Prioridad",
    "N.º de ítems",
    "Total estimado",
    "Total parcial",
    "Fecha requerida",
    "Justificación",
)

NOMBRE_HOJA = "Requerimientos"


def _fila_de(requerimiento):
    return (
        requerimiento.consecutivo,
        requerimiento.get_estado_display(),
        requerimiento.fecha_solicitud,
        requerimiento.solicitante,
        str(requerimiento.area),
        str(requerimiento.centro_costo),
        str(requerimiento.prioridad),
        requerimiento.items.count(),
        float(requerimiento.total_estimado),
        "Sí" if requerimiento.tiene_items_sin_precio else "No",
        requerimiento.fecha_requerida,
        requerimiento.justificacion,
    )


def generar_excel_requerimientos(requerimientos):
    """Construye el libro de Excel del consolidado de requerimientos.

    Recibe el queryset ya filtrado (los mismos criterios de la bandeja) y
    devuelve un `Workbook` de `openpyxl`, listo para guardarse en la
    respuesta HTTP. Las columnas de fecha quedan como `date` real, no texto:
    así quien lo reciba puede seguir filtrando y calculando en Excel sin
    tener que volver a convertir nada.
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = NOMBRE_HOJA

    hoja.append(ENCABEZADOS)
    for celda in hoja[1]:
        celda.font = Font(bold=True)

    fila = 1
    for requerimiento in requerimientos.select_related(
        "area", "centro_costo", "prioridad"
    ).prefetch_related("items"):
        fila += 1
        hoja.append(_fila_de(requerimiento))
        hoja.cell(row=fila, column=3).number_format = "yyyy-mm-dd"
        hoja.cell(row=fila, column=11).number_format = "yyyy-mm-dd"
        hoja.cell(row=fila, column=9).number_format = "#,##0.00"

    # Tabla con filtros nativos de Excel: el analista puede seguir acotando
    # el archivo sin tener que seleccionar el rango a mano.
    if fila > 1:
        rango = f"A1:{get_column_letter(len(ENCABEZADOS))}{fila}"
        tabla = Table(displayName="Requerimientos", ref=rango)
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        hoja.add_table(tabla)

    anchos = (16, 14, 16, 22, 20, 20, 10, 10, 14, 12, 14, 40)
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    return libro
